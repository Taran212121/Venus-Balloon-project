"""File outputs: JSON, CSV, mission tables, Excel workbook, LaTeX macros.

Everything here is plain stdlib except the workbook (openpyxl, optional --
skipped with a note when not installed).
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from .models import Config, Report, Summary, Trajectory, to_plain


# --------------------------------------------------------------------------
# Generic writers
# --------------------------------------------------------------------------


def write_json(path: str | Path, payload: Any) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, default=_jsonable) + "\n", encoding="utf-8")
    return path


def write_csv(path: str | Path, header: list[str], rows: list[list[Any]]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
    return path


def _jsonable(value: Any) -> Any:
    if isinstance(value, tuple):
        return list(value)
    if hasattr(value, "to_plain"):
        return value.to_plain()
    if hasattr(value, "__dataclass_fields__"):
        return to_plain(value)
    return str(value)


# --------------------------------------------------------------------------
# Trajectory and ensemble tables
# --------------------------------------------------------------------------


def trajectory_rows(trajectory: Trajectory) -> tuple[list[str], list[list[float]]]:
    header = [
        "time_s",
        "altitude_m",
        "velocity_m_s",
        "mach",
        "deceleration_g",
        "dynamic_pressure_pa",
        "heat_flux_conv_w_m2",
        "heat_flux_rad_w_m2",
        "heat_flux_total_w_m2",
        "heat_load_conv_j_m2",
        "heat_load_rad_j_m2",
        "heat_load_total_j_m2",
        "payload_temperature_k",
        "payload_absorbed_energy_j",
        "density_kg_m3",
        "temperature_k",
    ]
    columns = [getattr(trajectory, name) for name in header]
    return header, [list(row) for row in zip(*columns)]


def ensemble_metric_rows(ens) -> tuple[list[str], list[list[Any]]]:
    header = [
        "run",
        "success",
        "failure_reason",
        "m_entry_kg",
        "fpa_deg",
        "rho_sigma",
        "peak_g",
        "peak_q_pa",
        "peak_heat_w_m2",
        "heat_load_j_m2",
        "payload_absorbed_energy_j",
        "mach_handoff",
        "handoff_altitude_m",
    ]
    rows = []
    for run in ens.trajectories:
        sample = run.input_sample
        row: list[Any] = [
            sample["label"],
            run.success,
            run.failure_reason or "",
            sample["m_entry_kg"],
            sample["fpa_deg"],
            sample["rho_sigma"],
        ]
        if run.success:
            row += [
                max(run.deceleration_g),
                max(run.dynamic_pressure_pa),
                max(run.heat_flux_total_w_m2),
                run.heat_load_total_j_m2[-1],
                run.payload_absorbed_energy_j[-1],
                run.mach[-1],
                run.altitude_m[-1],
            ]
        else:
            row += [""] * 7
        rows.append(row)
    return header, rows


def event_time_rows(summary: Summary) -> tuple[list[str], list[list[Any]]]:
    names = list(summary.event_time_arrays_s)
    n_runs = max((len(values) for values in summary.event_time_arrays_s.values()), default=0)
    rows = [
        [idx] + [summary.event_time_arrays_s[name][idx] for name in names] for idx in range(n_runs)
    ]
    return ["run_index"] + names, rows


def corridor_rows(summary: Summary) -> tuple[list[str], list[list[float]]]:
    corridor = summary.handoff_corridor
    altitudes = corridor.get("altitude_m", ())
    header = ["altitude_m"]
    columns: list[tuple[float, ...]] = [tuple(altitudes)]
    for metric, bands in corridor.items():
        if metric == "altitude_m":
            continue
        for band in ("p2_5", "p50", "p97_5"):
            header.append(f"{metric}_{band}")
            columns.append(tuple(bands[band]))
    return header, [list(row) for row in zip(*columns)]


# --------------------------------------------------------------------------
# Mission phase table
# --------------------------------------------------------------------------

_PHASE_LABELS = (
    ("T0_interface", "Entry interface (T0)"),
    ("peak_g", "Peak deceleration"),
    ("peak_q", "Peak dynamic pressure"),
    ("peak_heat", "Peak heat flux"),
    ("drogue_deploy", "Drogue / DGB deploy"),
    ("main_deploy", "Main parachute deploy"),
    ("handoff", "Hand-off to descent stage"),
)


def mission_phase_rows(trajectory: Trajectory) -> tuple[list[str], list[list[Any]]]:
    header = ["event", "time_s", "altitude_m", "velocity_m_s", "mach"]
    rows = []
    for key, label in _PHASE_LABELS:
        t = trajectory.event_times_s.get(key)
        if t is None:
            continue
        state = _state_at(trajectory, t)
        rows.append([label, round(t, 2), round(state[0], 1), round(state[1], 1), round(state[2], 3)])
    return header, rows


def mission_phase_markdown(trajectory: Trajectory) -> str:
    header, rows = mission_phase_rows(trajectory)
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines += ["| " + " | ".join(str(cell) for cell in row) + " |" for row in rows]
    return "\n".join(lines) + "\n"


def _state_at(trajectory: Trajectory, t_s: float) -> tuple[float, float, float]:
    times = trajectory.time_s
    if t_s <= times[0]:
        idx = 0
    elif t_s >= times[-1]:
        idx = len(times) - 1
    else:
        idx = min(range(len(times)), key=lambda i: abs(times[i] - t_s))
    return trajectory.altitude_m[idx], trajectory.velocity_m_s[idx], trajectory.mach[idx]


# --------------------------------------------------------------------------
# Workbook and LaTeX macros
# --------------------------------------------------------------------------


def write_workbook(path: str | Path, cfg: Config, summary: Summary, report: Report) -> Path | None:
    """Write the metrics workbook; returns None (with no side effects) if openpyxl is absent."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return None

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()

    metrics = book.active
    metrics.title = "metrics"
    metrics.append(["metric", "p2_5", "p50", "p97_5", "units", "source_id"])
    for name, stats in summary.metric_intervals.items():
        metrics.append([name, stats["p2_5"], stats["p50"], stats["p97_5"], stats["units"], stats["source_id"]])

    sobol = book.create_sheet("sobol_s1")
    inputs = sorted({key for values in summary.sobol_s1.values() for key in values})
    sobol.append(["metric", *inputs, "sum_S1", "interaction_or_residual", "dominant_driver"])
    for metric, values in summary.sobol_s1.items():
        sobol.append(
            [metric]
            + [values.get(name, 0.0) for name in inputs]
            + [
                summary.sobol_sum_s1[metric],
                summary.sobol_interaction_or_residual_share[metric],
                summary.dominant_driver[metric],
            ]
        )

    checks = book.create_sheet("validation")
    checks.append(["check", "passed", "detail"])
    for check in report.checks:
        detail = "; ".join(
            f"{key}={value}" for key, value in check.items() if key not in ("name", "passed", "violations")
        )
        checks.append([check["name"], check["passed"], detail])

    run_sheet = book.create_sheet("run")
    run_sheet.append(["key", "value"])
    for key, value in (
        ("tuple_id", cfg.architecture.get("tuple_id", "")),
        ("workbook_id", cfg.architecture.get("workbook_id", "")),
        ("n_requested", summary.n_requested),
        ("n_success", summary.n_success),
        ("n_failed", summary.n_failed),
        ("entry_fpa_deg", cfg.entry_fpa_deg),
        ("m_entry_kg", cfg.m_entry_kg),
        ("ballistic_coefficient_kg_m2", cfg.ballistic_coefficient_kg_m2),
    ):
        run_sheet.append([key, value])

    book.save(path)
    return path


def write_latex_macros(path: str | Path, cfg: Config, summary: Summary) -> Path:
    """\\newcommand macros for the report: p50 and p97_5 of each metric."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "% entry_traj auto-generated metric macros",
        f"% tuple_id: {cfg.architecture.get('tuple_id', '')}",
    ]
    for name, stats in summary.metric_intervals.items():
        macro = _macro_name(name)
        p50 = format(stats["p50"], ".6g")
        p975 = format(stats["p97_5"], ".6g")
        lines.append("\\newcommand{\\" + macro + "Pfifty}{" + p50 + "}")
        lines.append("\\newcommand{\\" + macro + "PninetysevenFive}{" + p975 + "}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _macro_name(metric: str) -> str:
    parts = [part for part in metric.replace("-", "_").split("_") if part]
    return "entrytraj" + "".join(part.capitalize() for part in parts)
